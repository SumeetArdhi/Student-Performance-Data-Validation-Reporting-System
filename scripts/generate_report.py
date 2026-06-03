"""
generate_report.py
------------------
Queries the SQLite database and produces a formatted Excel report with:
  Sheet 1 - Summary Dashboard   (KPIs + subject-wise table)
  Sheet 2 - All Student Records (colour-coded marks)
  Sheet 3 - Validation Issues   (audit log)
  Sheet 4 - At-Risk Students
"""

import sqlite3
import os
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# -- Paths -------------------------------------------------------------------
BASE_DIR   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH    = os.path.join(BASE_DIR, "db", "student_performance.db")
ISSUES_CSV = os.path.join(BASE_DIR, "data", "validation_issues.csv")
REPORT_OUT = os.path.join(BASE_DIR, "reports", "student_performance_report.xlsx")

# -- Colour palette ----------------------------------------------------------
C_HEADER_BG  = "1F4E79"
C_HEADER_FG  = "FFFFFF"
C_ALT_ROW    = "D6E4F0"
C_KPI_BG     = "E8F4FD"
C_WARN_BG    = "FFF2CC"
C_DANGER_BG  = "FCE4D6"
C_GREEN_BG   = "E2EFDA"
C_BORDER     = "BDD7EE"

_thin   = Side(style="thin", color=C_BORDER)
BORDER  = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

TITLE_FONT  = Font(name="Arial", bold=True, size=14, color=C_HEADER_BG)
HEADER_FONT = Font(name="Arial", bold=True, size=10, color=C_HEADER_FG)
BODY_FONT   = Font(name="Arial", size=10)
KPI_FONT    = Font(name="Arial", bold=True, size=18, color=C_HEADER_BG)
KPI_LBL     = Font(name="Arial", size=9,   color="595959")


def _header_fill(color=C_HEADER_BG):
    return PatternFill("solid", fgColor=color)


def _style_header_row(ws, col_count, bg=C_HEADER_BG):
    """Apply header styling to the current last row of ws."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.font      = HEADER_FONT
        cell.fill      = _header_fill(bg)
        cell.border    = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_col_widths(ws, df, start_col=1):
    """Set column widths based on content."""
    for i, col_name in enumerate(df.columns):
        col_ltr  = get_column_letter(start_col + i)
        max_data = df.iloc[:, i].astype(str).str.len().max() if len(df) else 0
        width    = min(max(len(str(col_name)), max_data) + 4, 40)
        ws.column_dimensions[col_ltr].width = width


def write_df_to_sheet(ws, df, alt_color=C_ALT_ROW):
    """Append a header row + data rows to ws with alternating row colours."""
    # Header
    ws.append(list(df.columns))
    _style_header_row(ws, len(df.columns))

    # Data rows
    for i, row in enumerate(df.itertuples(index=False), start=1):
        ws.append(list(row))
        bg = alt_color if i % 2 == 0 else "FFFFFF"
        for col in range(1, len(df.columns) + 1):
            cell           = ws.cell(row=ws.max_row, column=col)
            cell.font      = BODY_FONT
            cell.fill      = PatternFill("solid", fgColor=bg)
            cell.border    = BORDER
            cell.alignment = Alignment(horizontal="left", vertical="center")

    _auto_col_widths(ws, df)


# -- Sheet builders ----------------------------------------------------------

def build_summary(wb, conn):
    ws       = wb.active
    ws.title = "Summary Dashboard"
    ws.sheet_view.showGridLines = False

    # Title row
    ws.merge_cells("A1:N1")
    tc            = ws["A1"]
    tc.value      = "Student Performance Dashboard"
    tc.font       = Font(name="Arial", bold=True, size=16, color=C_HEADER_BG)
    tc.alignment  = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 10   # visual spacer

    # KPI boxes (each 2 cols wide, 2 rows tall, 1-col gap between)
    kpis = {
        "Total Students":     "SELECT COUNT(DISTINCT student_id) FROM students",
        "Avg Marks":          "SELECT ROUND(AVG(marks),1) FROM student_performance WHERE marks IS NOT NULL",
        "Avg Attendance (%)": "SELECT ROUND(AVG(attendance_pct),1) FROM student_performance",
        "Pass Rate (%)":      ("SELECT ROUND(SUM(CASE WHEN marks>=50 THEN 1 ELSE 0 END)"
                               "*100.0/COUNT(*),1) FROM student_performance WHERE marks IS NOT NULL"),
        "At-Risk Students":   "SELECT COUNT(*) FROM v_at_risk_students",
        "Top Performers":     "SELECT COUNT(*) FROM v_top_performers",
    }

    kpi_row = 3
    kpi_col = 1
    for label, sql in kpis.items():
        val = conn.execute(sql).fetchone()[0]
        r, c = kpi_row, kpi_col
        ws.merge_cells(start_row=r,   start_column=c, end_row=r,   end_column=c + 1)
        ws.merge_cells(start_row=r+1, start_column=c, end_row=r+1, end_column=c + 1)
        lbl_cell            = ws.cell(row=r,   column=c, value=label)
        val_cell            = ws.cell(row=r+1, column=c, value=val)
        lbl_cell.font       = KPI_LBL
        val_cell.font       = KPI_FONT
        lbl_cell.alignment  = Alignment(horizontal="center")
        val_cell.alignment  = Alignment(horizontal="center")
        for ri in (r, r + 1):
            for ci in range(c, c + 2):
                ws.cell(ri, ci).fill   = PatternFill("solid", fgColor=C_KPI_BG)
                ws.cell(ri, ci).border = BORDER
        ws.row_dimensions[r].height   = 18
        ws.row_dimensions[r+1].height = 28
        kpi_col += 3
        if kpi_col > 13:
            kpi_col  = 1
            kpi_row += 4

    # Subject-wise summary table
    subj_df     = pd.read_sql("SELECT * FROM v_subject_summary", conn)
    table_row   = kpi_row + 4

    # Section title — write directly into the correct row, then advance
    title_cell            = ws.cell(row=table_row, column=1, value="Subject-wise Performance Summary")
    title_cell.font       = TITLE_FONT
    ws.row_dimensions[table_row].height = 24
    table_row += 1

    # Write subject header at exact table_row using direct cell writes (not ws.append)
    for col_idx, col_name in enumerate(subj_df.columns, start=1):
        cell           = ws.cell(row=table_row, column=col_idx, value=col_name)
        cell.font      = HEADER_FONT
        cell.fill      = _header_fill()
        cell.border    = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    table_row += 1

    for i, row in enumerate(subj_df.itertuples(index=False), start=1):
        bg = C_ALT_ROW if i % 2 == 0 else "FFFFFF"
        for col_idx, val in enumerate(row, start=1):
            cell           = ws.cell(row=table_row, column=col_idx, value=val)
            cell.font      = BODY_FONT
            cell.fill      = PatternFill("solid", fgColor=bg)
            cell.border    = BORDER
            cell.alignment = Alignment(horizontal="center")
        table_row += 1

    for col_idx in range(1, len(subj_df.columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22


def build_all_students(wb, conn):
    ws       = wb.create_sheet("All Student Records")
    ws.sheet_view.showGridLines = False

    df = pd.read_sql("""
        SELECT s.student_id, s.name, s.age, s.gender, p.subject,
               p.marks, p.grade, p.attendance_pct, s.semester, s.email
        FROM students s
        JOIN student_performance p ON s.student_id = p.student_id
        ORDER BY s.student_id
    """, conn)

    # Title
    col_count = len(df.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    tc            = ws["A1"]
    tc.value      = "All Student Records"
    tc.font       = TITLE_FONT
    tc.alignment  = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.append([])   # blank spacer row

    write_df_to_sheet(ws, df)

    # Colour-code marks column
    marks_col = list(df.columns).index("marks") + 1
    data_start = 4   # row 1=title, row 2=blank, row 3=header, rows 4+ = data
    for row_idx in range(data_start, ws.max_row + 1):
        cell = ws.cell(row=row_idx, column=marks_col)
        try:
            val = float(cell.value)
            if val >= 75:
                cell.fill = PatternFill("solid", fgColor=C_GREEN_BG)
            elif val < 50:
                cell.fill = PatternFill("solid", fgColor=C_DANGER_BG)
        except (TypeError, ValueError):
            pass


def build_validation_issues(wb):
    ws       = wb.create_sheet("Validation Issues")
    ws.sheet_view.showGridLines = False

    if not os.path.exists(ISSUES_CSV):
        ws["A1"] = "No validation issues file found. Run validate_data.py first."
        return

    df         = pd.read_csv(ISSUES_CSV)
    col_count  = len(df.columns)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    tc            = ws["A1"]
    tc.value      = "Data Validation Issues Log"
    tc.font       = TITLE_FONT
    tc.alignment  = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.append([])   # blank spacer

    write_df_to_sheet(ws, df, alt_color=C_WARN_BG)


def build_at_risk(wb, conn):
    ws       = wb.create_sheet("At-Risk Students")
    ws.sheet_view.showGridLines = False

    df        = pd.read_sql("SELECT * FROM v_at_risk_students ORDER BY risk_category, marks", conn)
    col_count = len(df.columns)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    tc            = ws["A1"]
    tc.value      = "At-Risk Students Report"
    tc.font       = TITLE_FONT
    tc.alignment  = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.append([])   # blank spacer

    write_df_to_sheet(ws, df, alt_color=C_DANGER_BG)


# -- Main --------------------------------------------------------------------

def main():
    print("=" * 60)
    print("  Report Generation")
    print("=" * 60)

    if not os.path.exists(DB_PATH):
        print(f"[ERROR] Database not found: {DB_PATH}")
        print("        Run transform_load.py first.")
        sys.exit(1)

    conn = sqlite3.connect(DB_PATH)
    wb   = Workbook()

    print("[REPORT] Building Summary Dashboard ...")
    build_summary(wb, conn)

    print("[REPORT] Building All Student Records ...")
    build_all_students(wb, conn)

    print("[REPORT] Building Validation Issues sheet ...")
    build_validation_issues(wb)

    print("[REPORT] Building At-Risk Students sheet ...")
    build_at_risk(wb, conn)

    conn.close()
    os.makedirs(os.path.dirname(REPORT_OUT), exist_ok=True)
    wb.save(REPORT_OUT)

    print(f"\n[OUTPUT] Report saved -> {REPORT_OUT}")
    print("=" * 60)


if __name__ == "__main__":
    main()
